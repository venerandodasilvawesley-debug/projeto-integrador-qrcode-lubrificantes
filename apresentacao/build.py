#!/usr/bin/env python3
"""Monta a apresentação do PI no modelo SENAI + Arauco (Projeto Sucuriú).

Uso: ~/.local/share/venv-pptx/bin/python build.py
Entrada: base.pptx (modelo com os slides de conteúdo já duplicados) e img/.
Saída:   Apresentacao_PI_QRCode_Lubrificantes.pptx
"""
import copy
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

AZUL = RGBColor(0x13, 0x62, 0xAD)
LARANJA = RGBColor(0xE7, 0x4F, 0x11)
ESCURO = RGBColor(0x1F, 0x2A, 0x37)
SUAVE = RGBColor(0x55, 0x60, 0x6E)
CARTAO = RGBColor(0xF1, 0xF4, 0xF8)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
AMARELO = RGBColor(0xF2, 0xB7, 0x05)
TITULO = "Impact"      # fonte de título do modelo
TEXTO = "Arial"
IMG = "img/"

prs = Presentation("base.pptx")
S = list(prs.slides)


def caixa(sl, x, y, w, h, paragrafos, tam=16, cor=ESCURO, fonte=TEXTO, negrito=False, alinh=PP_ALIGN.LEFT,
          ancora=MSO_ANCHOR.TOP, espaco=6, nome=None):
    """paragrafos: str, ou lista de str / lista de [(texto, {opções}), ...] por parágrafo."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if nome:
        tb.name = nome
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    if isinstance(paragrafos, str):
        paragrafos = [paragrafos]
    for i, par in enumerate(paragrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinh
        p.space_after = Pt(espaco)
        runs = par if isinstance(par, list) else [(par, {})]
        for texto, op in runs:
            r = p.add_run()
            r.text = texto
            f = r.font
            f.name = op.get("fonte", fonte)
            f.size = Pt(op.get("tam", tam))
            f.bold = op.get("negrito", negrito)
            f.italic = op.get("italico", False)
            f.color.rgb = op.get("cor", cor)
    return tb


def titulo(sl, texto, sub=None):
    caixa(sl, 0.6, 1.0, 11.2, 0.75, texto, tam=34, cor=AZUL, fonte=TITULO, nome="Título")
    if sub:
        caixa(sl, 0.62, 1.72, 11.0, 0.4, sub, tam=16, cor=SUAVE, nome="Subtítulo")


def forma(sl, tipo, x, y, w, h, cor, raio=0.08):
    s = sl.shapes.add_shape(tipo, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = cor
    s.line.fill.background()
    s.shadow.inherit = False
    if tipo == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = raio
    return s


def cartao(sl, x, y, w, h, cor=CARTAO):
    return forma(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, cor)


def bolinha(sl, x, y, d, texto, cor=LARANJA, tam=20):
    c = forma(sl, MSO_SHAPE.OVAL, x, y, d, d, cor)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = texto
    r.font.size = Pt(tam)
    r.font.bold = True
    r.font.name = TEXTO
    r.font.color.rgb = BRANCO
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return c


def foto(sl, arq, x, y, w=None, h=None, nome=None):
    p = sl.shapes.add_picture(IMG + arq, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if nome:
        p.name = nome
    return p


def foto_cortada(sl, arq, x, y, w, h, nome=None):
    """Imagem preenchendo a caixa (como object-fit: cover), cortando o excesso."""
    from PIL import Image
    iw, ih = Image.open(IMG + arq).size
    p = sl.shapes.add_picture(IMG + arq, Inches(x), Inches(y), Inches(w), Inches(h))
    alvo, real = w / h, iw / ih
    if real > alvo:
        corte = (1 - alvo / real) / 2
        p.crop_left = p.crop_right = corte
    else:
        corte = (1 - real / alvo) / 2
        p.crop_top = p.crop_bottom = corte
    if nome:
        p.name = nome
    return p


def tabela(sl, x, y, w, linhas, larguras, tam=14, altura_linha=0.42):
    nl, nc = len(linhas), len(linhas[0])
    t = sl.shapes.add_table(nl, nc, Inches(x), Inches(y), Inches(w), Inches(altura_linha * nl)).table
    for j, lw in enumerate(larguras):
        t.columns[j].width = Inches(lw)
    for i, linha in enumerate(linhas):
        for j, val in enumerate(linha):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.05)
            c.fill.solid()
            c.fill.fore_color.rgb = AZUL if i == 0 else (CARTAO if i % 2 else BRANCO)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.size = Pt(tam)
            r.font.name = TEXTO
            r.font.bold = i == 0
            r.font.color.rgb = BRANCO if i == 0 else ESCURO
    return t


def notas(sl, texto):
    sl.notes_slide.notes_text_frame.text = texto


def capa(sl, t1, t2, extra):
    for sh in sl.shapes:
        if sh.name == "CaixaDeTexto 1":
            sh.left, sh.top, sh.width = Inches(0.88), Inches(2.35), Inches(6.6)
            r = sh.text_frame.paragraphs[0].runs[0]
            r.text = t1
            r.font.size = Pt(54)
        elif sh.name == "CaixaDeTexto 2":
            sh.left, sh.top, sh.width = Inches(0.98), Inches(4.15), Inches(6.4)
            r = sh.text_frame.paragraphs[0].runs[0]
            r.text = t2
            r.font.size = Pt(20)
    caixa(sl, 0.98, 5.15, 6.3, 1.1, extra, tam=14, cor=BRANCO, espaco=4)


# ---------------- 1. capa ----------------
capa(S[0], "QR CODE NA LUBRIFICAÇÃO",
     "Sistema digital de identificação e gestão de lubrificantes em um redutor de velocidade",
     ["Karine Dionisio Corrêa · Leonardo Trajano · Wesley Venerando da Silva",
      "Técnico em Mecânica · Orientador(a): [nome do orientador] · 2026"])
notas(S[0], "Apresentar a equipe e o tema: um QR Code colado no equipamento abre, no celular, tudo o que o mecânico "
            "precisa para lubrificar com o óleo certo, no ponto certo, e registrar o que fez.")

# ---------------- 2. contexto ----------------
sl = S[1]
titulo(sl, "Três Lagoas: a Capital Nacional da Celulose")
stats = [("2009", "Fibria (atual Suzano)", "1,3 milhão t/ano de celulose"),
         ("2012", "Eldorado Brasil", "1,7 milhão t/ano de celulose"),
         ("2021", "Lei nº 14.142", "título de Capital Nacional da Celulose")]
for i, (n, a, b) in enumerate(stats):
    x = 0.6 + i * 3.75
    cartao(sl, x, 2.05, 3.45, 2.1)
    caixa(sl, x + 0.3, 2.25, 2.9, 0.8, n, tam=44, cor=LARANJA, fonte=TITULO)
    caixa(sl, x + 0.3, 3.1, 2.9, 1.0, [[(a, {"negrito": True})], b], tam=15, espaco=2)
caixa(sl, 0.6, 4.55, 10.3, 1.5,
      [[("As máquinas principais já nasceram com lubrificação automática. ", {"negrito": True}),
        ("Mas redutores, bombas e transportadores continuam sendo lubrificados pelo mecânico, "
         "e é aí que este projeto atua.", {})]], tam=18)
notas(S[1], "Contexto regional. Fonte: introdução do documento do PI (referência 5). "
            "Ligar com as fábricas de celulose da região, inclusive o Projeto Sucuriú da Arauco.")

# ---------------- 3. problema ----------------
sl = S[2]
titulo(sl, "O problema", "Na hora de lubrificar, a informação não está junto do equipamento")
itens = [("Tempo perdido", "O mecânico procura manual, ficha ou supervisor para saber óleo, volume e pontos."),
         ("Risco de erro", "Óleo errado ou misturado: o sintético PAO não pode levar mineral nem poliglicol."),
         ("Registro falho", "Anotado em papel e depois, o registro se perde e o PCM fica sem histórico.")]
for i, (a, b) in enumerate(itens):
    x = 0.6 + i * 3.75
    cartao(sl, x, 2.45, 3.45, 3.0)
    bolinha(sl, x + 0.3, 2.7, 0.7, str(i + 1))
    caixa(sl, x + 0.3, 3.6, 2.9, 0.5, a, tam=22, cor=AZUL, fonte=TITULO)
    caixa(sl, x + 0.3, 4.15, 2.9, 1.2, b, tam=15, cor=ESCURO)
notas(sl, "Problema de pesquisa: de que forma um QR Code fixado no equipamento pode reduzir o tempo gasto pelo "
          "mecânico para identificar e registrar a lubrificação de um redutor de velocidade?")

# ---------------- 4. objetivo ----------------
sl = S[3]
titulo(sl, "Objetivo")
caixa(sl, 0.6, 1.9, 11.0, 1.2,
      "Desenvolver e avaliar um sistema com QR Code que leve ao mecânico, no próprio equipamento, "
      "as informações de lubrificação, reduzindo o tempo de identificação e de registro.", tam=22, cor=ESCURO)
passos = [("Levantar", "dados do redutor na placa, no manual e no equipamento"),
          ("Desenvolver", "a ficha digital e a etiqueta QR resistente"),
          ("Medir", "o tempo nos dois métodos: tradicional × QR Code"),
          ("Calcular", "a redução de tempo e o retorno financeiro")]
for i, (a, b) in enumerate(passos):
    x = 0.6 + i * 2.8
    bolinha(sl, x, 3.55, 0.65, str(i + 1), cor=AZUL)
    caixa(sl, x, 4.35, 2.5, 0.45, a, tam=20, cor=LARANJA, fonte=TITULO)
    caixa(sl, x, 4.85, 2.5, 1.1, b, tam=15)
notas(sl, "Pesquisa aplicada, quantitativa, em ambiente didático, comparando dois métodos de trabalho.")

# ---------------- 5. equipamento ----------------
sl = S[4]
titulo(sl, "O equipamento estudado", "Redutor coroa e rosca sem fim WEG-Cestari da oficina do SENAI Três Lagoas")
tabela(sl, 0.6, 2.3, 6.9, [
    ["Item", "Dado"],
    ["Linha", "Magma M, tamanho 06 (código M0612802V0D001)"],
    ["Redução", "1:25 · 1.750 → 70 rpm · 1,75 kW"],
    ["Engrenamento", "Rosca sem fim de aço, coroa de bronze"],
    ["Carcaça", "Ferro fundido · 18,1 kg (sem óleo)"],
    ["Lubrificação", "Imersão em óleo sintético PAO ISO VG 460"],
    ["Posição", "Vertical, placa de identificação para cima"],
], [1.9, 5.0], tam=14)
foto_cortada(sl, "image3.jpeg", 8.0, 2.3, 3.2, 3.95, nome="Foto medição")
caixa(sl, 8.0, 6.3, 3.2, 0.3, "Medição da carcaça com paquímetro", tam=11, cor=SUAVE)
notas(sl, "Placa: série 818100, fabricado em 2018. A linha Magma M tamanho 06 foi deduzida pelo código e pelo manual, "
          "sem confirmação do fabricante. Aplicação hipotética: acionamento de uma pequena esteira transportadora.")

# ---------------- 6. pontos ----------------
sl = S[5]
titulo(sl, "Pontos de lubrificação", "Identificados e fotografados no próprio redutor")
foto_cortada(sl, "red-nivel-dreno.jpg", 0.6, 2.25, 3.1, 4.1, nome="Foto pontos")
pts = [("Enchimento", "Pelo respiro de latão, na face de cima"),
       ("Nível", "Bujão do meio da face da tampa cega; não há visor"),
       ("Dreno", "Bujão mais baixo dessa face e bujão da base")]
for i, (a, b) in enumerate(pts):
    y = 2.35 + i * 1.3
    bolinha(sl, 4.2, y, 0.6, str(i + 1), cor=AZUL, tam=18)
    caixa(sl, 5.0, y - 0.05, 6.0, 0.45, a, tam=22, cor=LARANJA, fonte=TITULO)
    caixa(sl, 5.0, y + 0.45, 6.0, 0.6, b, tam=16)
notas(sl, "O nível é conferido retirando o bujão, com o equipamento parado e bloqueado. "
          "Na posição vertical adotada, o bujão de cima não é usado.")

# ---------------- 7. dados principais ----------------
sl = S[6]
titulo(sl, "O que o mecânico precisa saber")
dados = [("VG 460", "Óleo sintético PAO (CLP HC), conforme a placa"),
         ("0,87 L", "Volume para todas as posições de trabalho"),
         ("15.000 h", "Troca com o óleo a 80 °C (10.000 h a 85 °C; 7.500 h a 90 °C)")]
for i, (n, b) in enumerate(dados):
    x = 0.6 + i * 3.75
    caixa(sl, x, 2.1, 3.6, 1.1, n, tam=54, cor=LARANJA, fonte=TITULO)
    caixa(sl, x, 3.3, 3.3, 1.0, b, tam=16)
cartao(sl, 0.6, 4.65, 10.5, 1.1)
caixa(sl, 0.9, 4.75, 10.0, 0.9,
      [[("Atenção: ", {"negrito": True, "cor": LARANJA}),
        ("não misturar com óleo mineral nem com sintético PG. Carcaça acima de 75 °C indica óleo acima do limite de 90 °C.", {})]],
      tam=16, ancora=MSO_ANCHOR.MIDDLE)
notas(sl, "Dados do manual WEG-Cestari da linha coroa e rosca. Óleos recomendados para VG 460: ExxonMobil SHC 634, "
          "Shell Omala S4 GX 460, Klübersynth EG4-460.")

# ---------------- 8. solução ----------------
sl = S[7]
titulo(sl, "A solução: uma etiqueta, uma ficha no celular")
caixa(sl, 0.6, 1.95, 4.2, 2.0,
      ["Página web: não precisa instalar aplicativo.",
       "Depois do primeiro acesso funciona sem internet.",
       "Etiqueta de papel colada em cima do redutor, com fita crepe: custo zero no protótipo."], tam=16, espaco=10)
foto(sl, "qr-lubrificacao-RED-001.png", 0.6, 4.1, w=1.7, nome="QR RED-001")
caixa(sl, 2.5, 4.35, 2.3, 1.2, [[("ficha-lubrificacao", {"negrito": True})], ".vercel.app/l/RED-001"], tam=13, cor=SUAVE, espaco=0)
for i, arq in enumerate(["image4.png", "image5.png", "image6.png"]):
    foto(sl, arq, 5.0 + i * 2.0, 1.95, h=3.95, nome=f"Tela {i + 1}")
notas(sl, "Telas: segurança e vida útil do óleo; pontos no modelo 3D com fotos reais; registro em poucos toques. "
          "Os registros ficam no aparelho e podem ser exportados em CSV.")

# ---------------- 9. funções ----------------
sl = S[8]
titulo(sl, "O que a ficha faz", "Pensada para o mecânico: nada para calcular ou adivinhar")
fun = [("Segurança primeiro", "LOTO, EPIs e alerta de mistura de óleos"),
       ("Vida útil do óleo", "Semáforo e data da próxima troca, sem digitar nada"),
       ("Galão conferido", "QR do galão: com o óleo errado, o registro não salva"),
       ("Troca passo a passo", "Uma etapa por tela, com fotos do próprio redutor"),
       ("Registro rápido", "Inspeção, complemento ou troca em poucos toques"),
       ("Proposta SAP", "Reserva de material e nota de manutenção simuladas")]
for i, (a, b) in enumerate(fun):
    x = 0.6 + (i % 3) * 3.6
    y = 2.35 + (i // 3) * 1.95
    cartao(sl, x, y, 3.35, 1.7)
    bolinha(sl, x + 0.25, y + 0.25, 0.5, str(i + 1), cor=LARANJA, tam=16)
    caixa(sl, x + 0.9, y + 0.27, 2.3, 0.5, a, tam=17, cor=AZUL, negrito=True)
    caixa(sl, x + 0.25, y + 0.9, 2.95, 0.75, b, tam=14)
notas(sl, "Troca simulada em 14/10/2026: próxima troca prevista para 30/06/2028. Tag do óleo RED-001-OLE-04, "
          "galão laranja (código e cor fictícios, de exemplo). A integração com o SAP é só uma proposta mostrada na tela.")

# ---------------- 10. modelo 3D ----------------
sl = S[9]
titulo(sl, "Modelo 3D do redutor", "Vista explodida e em corte, com a ficha e o tag de cada peça")
foto(sl, "image7.png", 0.6, 2.3, h=3.95, nome="Vista explodida")
foto(sl, "image8.png", 6.05, 2.3, h=3.95, nome="Vista em corte")
notas(sl, "Tocando no número ou na peça, a ficha mostra nome, material, função, o que inspecionar e falhas comuns, "
          "além do tag (ex.: RED-001-COR-01). Peças internas são ilustrativas e devem ser confirmadas numa desmontagem.")

# ---------------- 11. cronometragem ----------------
sl = S[10]
titulo(sl, "Cronometragem", "Cinco participantes, os dois métodos, junto ao redutor da oficina")
tabela(sl, 0.6, 2.35, 10.4, [
    ["Etapa", "Método tradicional", "Método com QR Code"],
    ["A. Lubrificante certo", "Placa e manual em PDF", "Escanear e ler a ficha"],
    ["B. Volume e pontos", "Procurar no manual e no redutor", "Já aparece na tela"],
    ["C. Registro", "Formulário de papel", "Poucos toques no celular"],
], [3.0, 3.7, 3.7], tam=16, altura_linha=0.55)
cartao(sl, 0.6, 4.85, 10.4, 1.2)
caixa(sl, 0.9, 4.95, 9.9, 1.0,
      [[("Participantes: ", {"negrito": True}),
        ("Felipe Araujo, Bruno Henrique de Souza, Maria Eduarda de Jesus, Simone Martins e Elis Dantas, "
         "na oficina de mecânica do SENAI Três Lagoas.", {})]], tam=15, ancora=MSO_ANCHOR.MIDDLE)
notas(sl, "Mesma ordem de serviço para todos; cronômetro do celular. A aplicação do óleo em si não foi cronometrada, "
          "porque é igual nos dois métodos. ATENÇÃO: o documento prevê ordem contrabalanceada, mas os participantes "
          "2 a 5 fizeram primeiro o tradicional; quem faz o tradicional primeiro já conhece o equipamento no segundo, "
          "o que pode aumentar a vantagem do QR Code. Citar isso como limitação.")

# ---------------- 12. resultados ----------------
sl = S[11]
titulo(sl, "Resultados", "Tempo total por participante, em minutos")
nomes = ["Felipe", "Bruno", "Maria Eduarda", "Simone", "Elis"]
trad = [503.27, 643.84, 726.26, 654.55, 646.28]
qr = [213.21, 180.08, 141.18, 178.34, 177.94]
cd = CategoryChartData()
cd.categories = nomes
cd.add_series("Tradicional", [round(v / 60, 1) for v in trad])
cd.add_series("QR Code", [round(v / 60, 1) for v in qr])
gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(2.2), Inches(6.9), Inches(4.2), cd)
gf.name = "Gráfico de tempos"
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(13)
ch.legend.font.name = TEXTO
pl = ch.plots[0]
pl.gap_width = 60
pl.has_data_labels = True
dl = pl.data_labels
dl.number_format = '0.0'
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(12)
dl.font.name = TEXTO
for serie, cor in zip(pl.series, (AZUL, LARANJA)):
    serie.format.fill.solid()
    serie.format.fill.fore_color.rgb = cor
ch.value_axis.has_major_gridlines = True
ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xE2, 0xE8)
ch.value_axis.tick_labels.font.size = Pt(12)
ch.value_axis.tick_labels.font.name = TEXTO
ch.value_axis.format.line.fill.background()
ch.category_axis.tick_labels.font.size = Pt(13)
ch.category_axis.tick_labels.font.name = TEXTO
mt, mq = sum(trad) / len(trad), sum(qr) / len(qr)


def mmss(s):
    return f"{int(s // 60)}:{int(round(s % 60)):02d}"


caixa(sl, 8.0, 2.25, 3.3, 1.2, f"−{round((1 - mq / mt) * 100)} %", tam=72, cor=LARANJA, fonte=TITULO)
caixa(sl, 8.0, 3.5, 3.3, 0.8, "de tempo com o QR Code (média)", tam=16)
caixa(sl, 8.0, 4.25, 3.3, 1.2, [[("Tradicional: ", {"negrito": True}), (mmss(mt) + " min", {})],
                                [("QR Code: ", {"negrito": True}), (mmss(mq) + " min", {})]], tam=18, espaco=4)
caixa(sl, 8.0, 5.25, 3.0, 0.7, "Média dos 5 participantes · oficina do SENAI Três Lagoas", tam=11, cor=SUAVE)
notas(sl, f"Médias dos 5 participantes: tradicional {mmss(mt)} min, QR Code {mmss(mq)} min. "
          "Todos foram mais rápidos com o QR Code. Medido o tempo total (sem separar as etapas A, B e C).")

# ---------------- 13. retorno ----------------
sl = S[12]
titulo(sl, "Retorno financeiro", "Ganho só no tempo do mecânico (aplicação hipotética, sem linha parada)")
ganho = mt - mq
anual_h = ganho * 12 / 3600
valor = anual_h * 14.32
dados = [(mmss(ganho) + " min", "poupados em cada intervenção"),
         (f"{anual_h:.1f} h".replace(".", ","), "por ano, com 12 inspeções mensais"),
         (f"R$ {valor:.0f}", "por ano e por equipamento (R$ 14,32/h do mecânico)")]
for i, (n, b) in enumerate(dados):
    x = 0.6 + i * 3.75
    cartao(sl, x, 2.35, 3.45, 2.2)
    caixa(sl, x + 0.3, 2.55, 2.9, 0.9, n, tam=40, cor=LARANJA, fonte=TITULO)
    caixa(sl, x + 0.3, 3.5, 2.9, 1.0, b, tam=15)
caixa(sl, 0.6, 4.95, 10.4, 1.2,
      [[("O ganho cresce com o número de equipamentos: ", {"negrito": True}),
        (f"com 100 equipamentos, cerca de R$ {valor * 100:,.0f} por ano".replace(",", ".") +
         ", contra o custo baixo de etiqueta e impressão.", {})]], tam=17)
notas(sl, "Hora do mecânico: R$ 3.150/mês (CBO 9113-05, Novo Caged via Portal Salário) ÷ 220 h = R$ 14,32, "
          "sem encargos: estimativa conservadora. Se a busca de informação acontece com a máquina parada, "
          "o ganho real inclui o custo da hora de produção.")

# ---------------- 14. análise do óleo ----------------
sl = S[13]
titulo(sl, "Análise do óleo no laboratório", "Ensaio de aprendizagem · viscosímetro Marte MVD-8 · prof. João Vitor de Moura Barbosa")
foto_cortada(sl, "tor-dreno.jpg", 0.6, 2.3, 2.45, 3.6, nome="Foto dreno")
foto_cortada(sl, "tor-ensaio-usado.jpg", 3.2, 2.3, 2.45, 3.6, nome="Foto ensaio")
caixa(sl, 0.6, 5.95, 5.05, 0.3, "Dreno do óleo usado (26/10) e ensaio no viscosímetro", tam=11, cor=SUAVE)
tabela(sl, 5.95, 2.3, 5.4, [
    ["", "Óleo novo", "Óleo usado"],
    ["Medido", "145,6\u00a0mPa·s (22,7\u00a0°C)", "91,0\u00a0mPa·s (27,3\u00a0°C)"],
    ["Na mesma temperatura", "≈\u00a0110\u00a0mPa·s", "91,0\u00a0mPa·s"],
    ["Estimado a 40\u00a0°C", "≈\u00a066\u00a0cSt", "≈\u00a056\u00a0cSt"],
], [1.6, 1.9, 1.9], tam=13, altura_linha=0.55)
selo = forma(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 5.95, 4.75, 5.0, 1.2, AMARELO, raio=0.15)
caixa(sl, 6.2, 4.83, 4.6, 1.05, [[("ATENÇÃO · viscosidade −17 %", {"negrito": True, "tam": 20})],
                                 "Óleo afinou e mostra sinais de contaminação (verde, turvo, depósito)."],
      tam=14, cor=ESCURO, ancora=MSO_ANCHOR.MIDDLE, espaco=2)
notas(sl, "Ensaio feito no laboratório do SENAI com o auxílio do prof. João Vitor de Moura Barbosa. Óleo novo: Ipiranga Ipitur HLP ISO 68 (70,9 cSt a 40 °C na ficha; medido ≈ 66 cSt, dentro da faixa do VG 68). As medições foram levadas à mesma temperatura "
          "pela curva do próprio óleo (equação de Walther, ASTM D341). Critério usual: até 10 % normal, 10 a 20 % atenção, "
          "acima de 20 % crítico. Suspeitas: fluido de corte, água (depósito alaranjado) ou reposição com óleo mais fino.")

# ---------------- 15. análise no QR ----------------
sl = S[14]
titulo(sl, "A análise do óleo dentro do QR Code")
foto(sl, "tela-tor-topo.jpg", 0.6, 1.9, h=4.4, nome="Tela TOR-001")
caixa(sl, 3.7, 1.95, 7.2, 2.6,
      [[("TOR-001 (torno): ", {"negrito": True, "cor": AZUL}),
        ("o QR mostra a situação do óleo com semáforo, o que fazer, as fotos do ensaio e o histórico. "
         "Quem medir de novo digita a viscosidade e a temperatura, e a ficha calcula a variação sozinha.", {})],
       [("RED-001 (redutor): ", {"negrito": True, "cor": AZUL}),
        ("o mesmo cartão entra como proposta: o óleo VG 460 ainda não foi analisado. "
         "A amostra será coletada na troca, porque o redutor tem só 0,87 L.", {})]], tam=16, espaco=12)
foto(sl, "qr-tor-001.png", 3.7, 4.55, w=1.5, nome="QR TOR-001")
caixa(sl, 5.4, 4.75, 4.0, 1.0, [[("Leia e veja ao vivo", {"negrito": True})], "ficha-lubrificacao.vercel.app/e/TOR-001"],
      tam=14, cor=SUAVE, espaco=2)
notas(sl, "Mostrar ao vivo no celular. Antes da banca, abrir a ficha uma vez com dados móveis: o Wi-Fi de visitantes "
          "do SENAI bloqueia o site, mas depois do primeiro acesso a ficha funciona sem internet.")

# ---------------- 16. valores Arauco ----------------
sl = S[15]
titulo(sl, "Alinhado aos valores da Arauco", "Código de Conduta do Projeto Sucuriú × o que o projeto entrega")
vals = [("Segurança", RGBColor(0xE8, 0x75, 0x1A), "LOTO, EPIs e alertas aparecem antes de qualquer intervenção"),
        ("Compromisso", RGBColor(0xC8, 0x50, 0x4A), "Cada intervenção fica registrada, com data, responsável e foto"),
        ("Excelência e inovação", RGBColor(0x2F, 0xA8, 0x94), "QR Code, modelo 3D e análise do óleo no celular"),
        ("Trabalho em equipe", RGBColor(0xD9, 0xA3, 0x00), "Três projetos integradores no mesmo site"),
        ("Bom cidadão", RGBColor(0x8F, 0xA8, 0x1E), "Óleo usado recolhido e destinado ao rerrefino (CONAMA)")]
for i, (v, cor, txt) in enumerate(vals):
    y = 2.3 + i * 0.82
    c = forma(sl, MSO_SHAPE.OVAL, 0.6, y, 0.6, 0.6, cor)
    caixa(sl, 1.4, y + 0.08, 3.7, 0.5, v, tam=20, cor=cor, fonte=TITULO)
    caixa(sl, 5.2, y + 0.12, 5.9, 0.6, txt, tam=16)
notas(sl, "Valores do Código de Conduta do Projeto Sucuriú (Arauco): Segurança, Compromisso, Excelência e inovação, "
          "Trabalho em equipe e Bom cidadão. Mostrar como cada um aparece na solução.")

# ---------------- 17. conclusão ----------------
sl = S[16]
titulo(sl, "Conclusão", "Informação certa no ponto de uso: menos tempo e menos erro")
for k, (cab, itens, cor) in enumerate([
    ("Contribuições", [f"Tempo cerca de {round((1 - mq / mt) * 100)} % menor com o QR Code",
                       "Acesso padronizado aos dados de lubrificação",
                       "Histórico confiável para o PCM e a TPM",
                       "Análise do óleo com semáforo no próprio QR"], AZUL),
    ("Próximos passos", ["Testar em linha de produção real",
                         "Analisar o óleo do redutor na próxima troca",
                         "Sincronizar os registros entre aparelhos",
                         "Integrar de fato ao SAP PM"], LARANJA)]):
    x = 0.6 + k * 5.4
    cartao(sl, x, 2.35, 5.1, 3.6)
    caixa(sl, x + 0.35, 2.55, 4.4, 0.5, cab, tam=24, cor=cor, fonte=TITULO)
    caixa(sl, x + 0.35, 3.2, 4.4, 2.6, ["• " + t for t in itens], tam=16, espaco=8)
notas(sl, "Limitações: cronometragem em ambiente didático; ordem não contrabalanceada nos participantes 2 a 5; "
          "linha Magma M 06 deduzida sem confirmação do fabricante; registros só no aparelho por enquanto.")

# ---------------- 18. obrigado ----------------
sl = S[17]
capa(sl, "OBRIGADO!", "Perguntas?", ["Leia os QR Codes e veja as fichas ao vivo"])
for i, (arq, rot) in enumerate([("qr-lubrificacao-RED-001.png", "RED-001 · lubrificação"), ("qr-tor-001.png", "TOR-001 · análise do óleo")]):
    x = 0.98 + i * 2.3
    forma(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, 5.6, 1.6, 1.6, BRANCO, raio=0.06)
    foto(sl, arq, x + 0.1, 5.7, w=1.4)
    caixa(sl, x, 7.22, 1.6, 0.25, rot.split(" · ")[0], tam=11, cor=BRANCO, alinh=PP_ALIGN.CENTER)
notas(sl, "Abrir para perguntas e, se possível, ler a etiqueta colada no redutor com o celular.")

prs.save("Apresentacao_PI_QRCode_Lubrificantes.pptx")
print("ok")
